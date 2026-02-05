import os
import datetime
import logging
import csv
import zipfile
import json
import geopandas as gpd
from  util import util_services
from pyogrio import list_layers
import re
import time
import pandas as pd
import oracledb
from util import std_descr
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment,PatternFill
from collections import defaultdict
from time import strftime
from time import gmtime
from datetime import datetime
import sqlite3
import io
import networkx as nx
import pandas as pd
from shapely import wkb
from shapely.geometry import MultiLineString
import numpy as np
class IQGeoCDIF:
    def __init__(self, output_folder_dir, folder_name,  config_json_file_path,input_gdb , extent="0", route_cdif_file="", route_iqgeo_file="",str_iqgeo_file=""):
        self.Initialized=False
        self.DataMigrationStatus=0
        self.folder_name=folder_name.lower()
        self.rootDirectory = output_folder_dir
        self.folder_path=os.path.join(output_folder_dir,folder_name)
        self.route_cdif_file = route_cdif_file
        self.config_json_file_path = config_json_file_path
        self.route_iqgeo_file = route_iqgeo_file
        self.str_iqgeo_file = str_iqgeo_file
        self.input_gdb = input_gdb 
        self.category= "CONDUIT"
        self._created_files = []

        if extent=="0":
            # use Saskatchewan (Canada) default extent if no value passed
            self.extent= "01030000000100000005000000624817016e815bc0e82d078e939f4b407a6431c4c3835bc0d8f547621b7c484076a008e6bf5659c008db8af2698248400e846c5e387e59c0d0b5d49167a64b40624817016e815bc0e82d078e939f4b40"
        else:
            self.extent = extent
            
        self._create_folder_for_cdif()
        self._create_design_metadata()

    def _create_folder_for_cdif(self):
        logging.info(f"Creating folder for CDIF | '{self.folder_path}'")
        try:
            if not os.path.exists(self.folder_path):
                os.makedirs(self.folder_path ) 
        except:
            logging.error(str(datetime.now())+" | "+"ERROR IN Creating folder | "+self.folder_path)
        return
    
    def _create_design_metadata(self):
        self.packagemetadata_file_name=self.rootDirectory +"\\package.metadata"
        logging.info(f"Creating Package metadata file | '{self.packagemetadata_file_name}'")
        
        data=[
            ["property","value"],
            ["format","cdif"],
            ["coord_system","4326"],
            ["boundary",self.extent]
        ]
        try:
            with open(self.packagemetadata_file_name,'w',newline='') as csvfile:
                writer=csv.writer(csvfile)
                writer.writerows(data)
            print(f"Package Metadata written sucessfully.")
            logging.info(f"Package metadata written sucessfully.'{self.packagemetadata_file_name}'")
        except Exception as e:
            print(f"Package metadata file could not be written. Error occured : '{str(e)}'")
            logging.error(f"Package metadata file '{self.packagemetadata_file_name}' could not be written. Error occured : '{str(e)}' ")
    def _create_field_mapping_file( self ):
        self.field_file_name= self.folder_path+"\\"+self.target_object.lower()+".fields"
        logging.info(f"Creating field mapping file  '{self.field_file_name}' for object {self.target_object} ")
        print(f"Creating field mapping file  '{self.field_file_name}' for object '{self.target_object}'")
        try:
            logging.info(f"Loading configuration JSON file  '{self.config_json_file_path}' ")
            fields=[]
            
            hdr_info=[["name","type","unit"]]
            fields.append(hdr_info)        
        
            hdr_info.append(["id","id",""]) #ID field column information 
            hdr_info.append([self.geomfld,self.geomtype,""])  

            for fld in self.reference_fields:
                hdr_info.append([fld['Field_Name'].lower(),fld['Field_Type'],])
                
            for fld in self.filtered_fields:
                hdr_info.append([fld['Field_Name'].lower(),fld['Field_Type'],])

                    # fields_issues.append([fld['Field_Name'],0,0,0,fld['Mandatory']])

            with open(self.field_file_name,'w',newline='') as csvfile:
                writer=csv.writer(csvfile)
                writer.writerows(hdr_info)
            
            print("Field Metadata written sucessfully for " + self.target_object)
            logging.info( f"Creating field mapping file completed for '{self.target_object}'")
            # else:
            #     print("ERROR: Geometry field not found in the config." )
            #     logging.error( f"ERROR:Geometry field not found in the config '{self.object_config_json}'. ")

        except Exception as e:
            print(f'Field metadata file could not be written. Error occured : {str(e)}')
            logging.error( f"Error occured in create field mapping file '{self.field_file_name}'. Error is {str(e)}  ")
        rel = os.path.relpath(self.field_file_name, self.rootDirectory)
        self._created_files.append(rel)

    def InitiateProcess(self ):
        self.start_time = time.time()
        with open(self.config_json_file_path, 'r') as file:
            self.config_json = json.load(file)

        self.records_skipped_due_to_issue=0
        self.records_skipped_due_to_parent_data_mismatch_issue=0
        self.number_of_records_truncated=0
        self.number_of_values_truncated=0

        in_structure_fc= ""
        in_structure_right_clm=""
        in_structure_left_clm=""
        
        out_structure_fc= ""
        out_structure_clm= ""
        out_structure_left_clm=""

        structure_fc= ""
        structure_clm= ""
        structure_left_clm=""        
        structure_right_clm=""

        self.target_object=self.config_json.get("Name").lower()
        self.src_lyr=self.config_json.get("NE_Table")
        self.ref_structure_lyr=self.config_json.get("NE_Structure_Layer")
        self.filter_column=self.config_json.get("Filter")
        
        self.mf = self.config_json.get("Attribute_Finalised")
        self.reference_fields = [flds for flds in self.mf if flds["FieldSource"].upper()=="IQGEO" and  flds["Field_Type"].lower()=="reference"  ]  #and flds["Join"] !=""
        for flds in self.reference_fields:
                if flds['Join']!='':
                    left_join_cond=flds['Join'].split('=')[0]
                    right_join_cond=flds['Join'].split('=')[1]
                    left_table=left_join_cond.split('.')[0]

                    if left_table!=self.src_lyr:
                        print(f"Warning!! Joining expression {flds['Join']} in referenced column referring to other then {self.src_lyr} layer. ")

                    left_col=left_join_cond.split('.')[1]
                    right_table=right_join_cond.split('.')[0]
                    right_col=right_join_cond.split('.')[1]

                if flds['Field_Name']=="in_structure":      # To handle ROUTE and STRUCTURE Referencing 
                    in_structure_fc= right_table.upper()
                    self.in_structure_fc=right_table.upper()
                    in_structure_right_clm= right_col.upper() #
                    in_structure_left_clm=left_col.upper()
                    self.in_structure_left_clm=in_structure_left_clm
                if flds['Field_Name']=="out_structure":  # To handle ROUTE and STRUCTURE Referencing 
                    out_structure_fc= right_table.upper()
                    self.out_structure_fc=right_table.upper()
                    out_structure_right_clm= right_col.upper()
                    out_structure_left_clm=left_col.upper()
                    self.out_structure_left_clm=out_structure_left_clm
                    #out_structure_gdf=gpd.read_file(self.input_gdb,layer=right_table,  columns=['object_id', right_col])
                
                 

        self.filtered_fields = [flds for flds in self.mf if flds["FieldSource"]=="NE"]
        self.geo_fields = [flds for flds in self.mf if (flds["Field_Type"]) =="Geometry"]
        NE_Fields_Name=[]

        for flds in self.filtered_fields:
                if flds['Source']=="":
                    NE_Fields_Name.append(flds['Field_Name'])
                else:
                    NE_Fields_Name.append(flds['Source'])
        
        self.NE_Fields= NE_Fields_Name
        if len(self.geo_fields)>0:    
            self.geomfld=self.geo_fields[0]["Field_Name"]
            self.geomtype= self.config_json.get("Geometry_Type")
        else:
            print("ERROR: Geometry field not found in the config." )
            logging.error( f"ERROR:Geometry field not found in the config '{self.config_json}'. ")
            return
        
        self._create_field_mapping_file()
        layers = list_layers(self.input_gdb)
        if self.src_lyr in layers:
            print(f"GDB Open Start at {time.time()}")
            self.org_gdf=gpd.read_file(self.input_gdb,layer=self.src_lyr )

            # self.org_gdf=gpd.read_file(self.input_gdb,layer=self.src_lyr,columns=NE_Fields_Name)

            print(f"GDB read completed at {time.time()}")
            
            print("Filtering for  " + self.filter_column) 
            mergedgdf=util_services.UTILServices.apply_sql_like_filter(self.org_gdf,self.filter_column) 

            if in_structure_fc!="": #In structure joining needs to be done
                instr_gdf=gpd.read_file(self.input_gdb,layer=right_table,  columns=['object_id', in_structure_right_clm,'TYPE_NAME'])
                instr_gdf= instr_gdf.rename(columns={'object_id': 'instr_object_id', in_structure_right_clm: 'INSTR_'+in_structure_right_clm,'TYPE_NAME': 'INSTR_TYPE_NAME'})   

                self.in_structure_right_clm='INSTR_'+in_structure_right_clm

                instr_gdf_atr=instr_gdf.drop(columns='geometry')
                instr_gdf_geom=instr_gdf[[self.in_structure_right_clm, 'geometry']].rename(columns={"geometry":"InStrGeom"}) 
                mergedgdf=mergedgdf.merge(instr_gdf_atr,  left_on=in_structure_left_clm, right_on=self.in_structure_right_clm, how='left' ).merge(instr_gdf_geom, left_on=in_structure_left_clm, right_on=self.in_structure_right_clm, how='left')
                
            if out_structure_fc!="": #Out structure joining needs to be done
                outstr_gdf=gpd.read_file(self.input_gdb,layer=right_table,  columns=['object_id', out_structure_right_clm,'TYPE_NAME'])
                outstr_gdf= outstr_gdf.rename(columns={'object_id': 'outstr_object_id', out_structure_right_clm: 'OUTSTR_'+out_structure_right_clm,'TYPE_NAME': 'OUTSTR_TYPE_NAME'})
                self.out_structure_right_clm='OUTSTR_'+out_structure_right_clm
                outstr_gdf_atr=outstr_gdf.drop(columns='geometry')
                outstr_gdf_geom=outstr_gdf[[self.out_structure_right_clm, 'geometry']].rename(columns={"geometry":"OutStrGeom"}) 
                mergedgdf=mergedgdf.merge(outstr_gdf_atr,  left_on=out_structure_left_clm, right_on=self.out_structure_right_clm, how='left' ).merge(outstr_gdf_geom, left_on=out_structure_left_clm, right_on=self.out_structure_right_clm, how='left')    
            
            if structure_fc!="": # Joinig to get Pole Anchor Structure reference
                str_gdf=gpd.read_file(self.input_gdb,layer=right_table,  columns=['object_id',structure_right_clm,'TYPE_NAME'])
                str_gdf= str_gdf.rename(columns={'object_id': 'str_object_id', structure_right_clm: 'STR_'+structure_right_clm,'TYPE_NAME': 'STR_TYPE_NAME'})
                self.structure_right_clm='STR_'+structure_right_clm
                str_gdf_atr=str_gdf.drop(columns='geometry')
                str_gdf_geom=str_gdf[[self.structure_right_clm, 'geometry']].rename(columns={"geometry":"StrGeom"}) 
                mergedgdf=mergedgdf.merge(str_gdf_atr,  left_on=structure_left_clm, right_on=self.structure_right_clm, how='left' )
                                #    .merge(str_gdf_geom, left_on=structure_left_clm, right_on=self.structure_right_clm, how='left')    
            
            
            
            self.gdf=mergedgdf
            self.total_records=len(self.gdf)
            if self.input_gdb.endswith(".gdb"):
                sqlite_db = self.input_gdb.replace(".gdb", ".sqlite")
                self.sqlite_db=sqlite_db
                if os.path.exists(sqlite_db):
                    print(f'SQL DB {sqlite_db} already exists, bypassing.')
                else:
                    print('Exporing selected layers in sqlite db')
                    util_services.UTILServices.export_gdb_to_sqlite(self.input_gdb,sqlite_db,['SPAN','SPAN_UNIT','SPAN_SPAN_UNIT','EQUIPMENT','SLOT','PLUGIN','PORT','SLOT_PLUGIN','CHASSIS','STRUCTURE_UNIT','TRANSMEDIA'])


            if self.total_records>0:
                self.Initialized=True
            else:
                print("No data found to be processed. Terminating the job.")
        else:
            print(f"Source  layer {self.src_lyr} is missing in FileGDB."  ) 

    def ProcessRecords(self):
        if self.Initialized==True:
            print("Starting data migration...")
            self.DataMigrationStatus=1
            rec_processed=0
            csv_data=[]
            # string_truncate_object=[]
            # duplicate_geom=[]
            # summary_rpt=[]
            # bypassed_data=[]
            # fields_issues=[]
            # type_mismatch_data=[]
            mandatory_objects_missing=[]

            self.string_truncate_object=[]
            self.bypassed_data=[]
            self.fields_issues= []
            self.type_mismatch_data= []
            self.parent_missing=[]
            self.topology_issues=[]
            self.skipped_conduit=[]
            self.core_hole_merged=[]
            
            clm_name=[]
            clm_name.append("id") #ID field column information 
            clm_name.append(self.geomfld)  
            
            for fld in self.reference_fields:
                self.fields_issues.append([fld['Field_Name'],0,0,0,fld['Mandatory'],0])
                clm_name.append(fld['Field_Name'].lower())

            for fld in self.filtered_fields:
                self.fields_issues.append([fld['Field_Name'],0,0,0,fld['Mandatory'],0])
                clm_name.append(fld['Field_Name'].lower())

            print("Processing data...  " ,end="")

            csv_data,rec_processed=self.ProcessCoundit(clm_name)

            print( str(rec_processed) + " rows processed." )
            flindx=-1
            for cl in clm_name:
                flindx+=1
                if cl.lower()=="name":
                    break             
            if flindx==-1:
                flindx=2
            number_of_geometry_duplicate,duplicate_geom=util_services.UTILServices.find_duplicates(csv_data,flindx)
            self.number_of_geometry_duplicate=number_of_geometry_duplicate
            self.duplicate_geom = duplicate_geom
            # self.string_truncate_object=string_truncate_object
            # self.bypassed_data=bypassed_data
            # self.fields_issues=fields_issues
            # self.type_mismatch_data=type_mismatch_data
            csv_data= util_services.UTILServices.remove_duplicates(csv_data)
            self.total_rec_processed =len( csv_data) #rec_processed
            self.csv_file_name=self.folder_path+ "\\"+self.target_object+".csv"
          
            self._WriteCSVFile(self.csv_file_name ,clm_name,csv_data)        
           
            self.DataMigrationStatus=2

        else:
            print("Proces is not initlizaed properly, either Initilization could not be performed OR some issue may occured during initialization.")

        return
    def GetColumnIndex(self,clm_names,clm):
        idx=0
          
        for cl in clm_names:
            if cl==clm:
                return idx
            idx+=1
        
        return -1
    
    
    
    def ProcessCoundit(self, clm_name):
        sqlite_path = self.rootDirectory +"\\"+util_services.UTILServices.generate_random_string(10)
         
        conn = sqlite3.connect(sqlite_path)
        conn2=sqlite3.connect(self.sqlite_db)
        cur = conn.cursor()
        cur2 = conn2.cursor()

        create_sql = f"CREATE TABLE IF NOT EXISTS route_geom (str_hash TEXT, geom TEXT, lngth REAL , no_vertex INTEGER , obj_id TEXT, from_str_id TEXt, to_str_id TEXT, from_str_type TEXt, to_str_type TEXT, span_type TEXT,in_str_name TEXT,out_str_name TEXT, is_active TEXT, cr_id TEXT);"
        cur.execute(create_sql)
        create_sql = f"CREATE INDEX idx_str_hash ON route_geom (str_hash);"
        cur.execute(create_sql)

        route_df = pd.read_csv(self.route_cdif_file) 
         
        csv_data=[]
        Missing_route_conduit=[]
        corehole_data=[]
        additional_conduit_run=[]
        max_id_conduit_run =0
        rec_processed=0
        total_records= len(self.gdf)
        src_crs=self.gdf.crs
        tgt_crs="EPSG:4326"
        obj_id_pre_fix=self.target_object #self.category #"CONDUIT" #self.src_lyr.replace(" ", "_")

        uqfld=self.config_json.get("Source_ID")  
        frm_str_fld_indx=util_services.UTILServices.GetColumnIndex(clm_name,'in_structure')
        to_str_fld_indx=util_services.UTILServices.GetColumnIndex(clm_name,'out_structure')
        root_housing_idx=util_services.UTILServices.GetColumnIndex(clm_name,'root_housing')
        housing_idx=util_services.UTILServices.GetColumnIndex(clm_name,'housing')
        conf_val_idx = util_services.UTILServices.GetColumnIndex(clm_name,'confidence_percentage')
        df_route_ref = pd.read_csv(self.route_iqgeo_file,  on_bad_lines='skip')#,nrows=10
        df_str_ref = pd.read_csv(self.str_iqgeo_file,  on_bad_lines='skip')#,nrows=10

        self.gdf.sort_values(by=['CALCULATED_LENGTH','SHAPE_Length'],  inplace=True)
        self.gen_conduit_id=0
        for index,row in self.gdf.iterrows():
            csv_row=[]
            confidence_val="0"
            trunc_record=False
            db_type_issue=False
            frm_str_obj_id=""
            to_str_obj_id=""
            housing_route_id=""
            # sql_to_check_no_of_conduits_in_span_asssociation=f"select count(1) no_of_asso from SPAN_SPAN_UNIT s where s.SPAN_UUID = '{row['UUID']}'" 
            # sql_to_check_no_of_conduits_in_span_asssociation=f"select su.* from SPAN_UNIT su inner join SPAN_SPAN_UNIT ssu  on su.UUID  = ssu.SPAN_UNIT_UUID where ssu.SPAN_UUID= '{row['UUID']}'" 
            sql_to_check_no_of_conduits_in_span_asssociation=f"select su.span_unit_name , su.type_name,su.span_unit_ref_name specification   ,su.inventory_status_code,su.work_order_name,su.account_code st_std_desc ,su.measured_length,su.ducts_available,su.diameter,su.label,su.st_rmks,su.st_vntge_yr,su.st_sap_ntwk_id  sap_pm_order ,sp.created_user,sp.created_date,sp.last_edited_user,sp.last_edited_date,su.uuid from SPAN_UNIT su inner join SPAN_SPAN_UNIT ssu  on su.UUID  = ssu.SPAN_UNIT_UUID inner join SPAN sp  on ssu.SPAN_UUID=sp.uuid where sp.uuid = '{row['UUID']}'" 

            cur2.execute(sql_to_check_no_of_conduits_in_span_asssociation )
            associated_span = cur2.fetchall()
            no_of_asso=len(associated_span)
            cur2.close
            span_nm= row['SPAN_NAME']
            # if span_nm=='OSP:COND::17054374':
            #     span_nm=='OSP:COND::17054374'
            # If multiple rows found in span unit association then similar number of conduits has to be added    
            if no_of_asso>1:
                print(f"Multiple Span Association found for {row['UUID']}")
                for asso_span_row in associated_span:
                    csv_row=[]
                    confidence_val="0"
                    # idval=re.sub(r'[^0-9a-zA-Z]', '', asso_span_row[17])
                    # iqgeo_ref_id_val= obj_id_pre_fix + idval
                    self.gen_conduit_id+=1
                    iqgeo_ref_id_val= obj_id_pre_fix + '/'+ str(self.gen_conduit_id) 

                    str_geom=row["geometry"]  #geomfld
                    str_geom=util_services.UTILServices.ConvertMultiLineToLine(str_geom)
                    csv_row.append(iqgeo_ref_id_val)
                    csv_row.append("")
                    fld_indx=0
                    for fld in self.reference_fields:
                        fld_val =""
                        if fld['Field_Name']=="in_structure":
                            str_name=row[self.in_structure_left_clm.upper()]
                            match = df_str_ref[df_str_ref["name"].astype(str).str.strip().str.upper() == str_name]
                            if not match.empty:
                                structure_id = str(match.iloc[0]["id"]).strip()
                                fld_val = structure_id
                                frm_str_obj_id = structure_id
                                # in_str_geom=row["InStrGeom"]
                            else:
                                fld_val = f"{str_name}"#str_name
                                # fld_val=str_name
                                frm_str_obj_id = ""
                                self.parent_missing.append([row[uqfld],"housing",row[self.structure_right_clm + "_x"],row[self.structure_right_clm + "_y"]])

                        # if fld['Field_Name']=="in_structure":
                        #     if pd.isnull(row["instr_object_id"]):
                        #         frm_str_obj_id="" 
                        #         self.parent_missing.append([row[uqfld],"in_structure", row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                        #     else:
                        #         frm_str_obj_id=re.sub(r'[^0-9a-zA-Z]', '', row["instr_object_id"])  

                        #         in_str_geom=row["InStrGeom"]
                        #         validated_obj=util_services.UTILServices.check_and_convert_value(fld_val,"str")
                        #         if validated_obj['is_null']==True or validated_obj['is_nan']  ==True:
                        #             frm_str_obj_id=""
                        #             self.parent_missing.append([row[uqfld],"in_structure", row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                        #         else:
                        #             frm_str_obj_id =self.in_structure_fc.upper()+str((frm_str_obj_id))
                        #     fld_val=frm_str_obj_id

                        if fld['Field_Name']=="out_structure":
                            str_name=row[self.out_structure_left_clm.upper()]
                            match = df_str_ref[df_str_ref["name"].astype(str).str.strip().str.upper() == str_name]
                            if not match.empty:
                                structure_id = str(match.iloc[0]["id"]).strip()
                                fld_val = structure_id
                                to_str_obj_id = structure_id
                                # out_str_geom=row["OutStrGeom"]
                            else:
                                fld_val = f"{str_name}"#str_name
                                # fld_val=str_name
                                to_str_obj_id = ""
                                self.parent_missing.append([row[uqfld],"housing",row[self.structure_right_clm + "_x"],row[self.structure_right_clm + "_y"]])


                            # if pd.isnull(row["outstr_object_id"]):
                            #     to_str_obj_id="" 
                            #     self.parent_missing.append([row[uqfld],"out_structure",  row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                            # else:
                            #     to_str_obj_id=re.sub(r'[^0-9a-zA-Z]', '', row["outstr_object_id"])   
                            #     out_str_geom=row["OutStrGeom"]

                            #     validated_obj=util_services.UTILServices.check_and_convert_value(fld_val,"str")
                            #     if validated_obj['is_null']==True or validated_obj['is_nan']  ==True:
                            #         to_str_obj_id=""
                            #         self.parent_missing.append([row[uqfld],"out_structure",  row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                            #     else:
                            #         to_str_obj_id =self.in_structure_fc.upper()+str((to_str_obj_id))
                            # fld_val=to_str_obj_id

                        csv_row.append( fld_val)
                        fld_indx+=1

                 # Read Route CSV for Housing reference
                    conduit_last_part=util_services.UTILServices.get_char_after(row['SPAN_NAME'],':')
                    conduit_last_numpart= conduit_last_part[1:]
                    combined_conditions  = f"((route_df['in_structure'] =='{frm_str_obj_id}') &  (route_df['out_structure'] ==  '{to_str_obj_id }'))  | ((route_df['in_structure'] =='{to_str_obj_id}') & (route_df['out_structure'] ==  '{frm_str_obj_id}'))"
                    filtered_df = route_df[eval(combined_conditions)]
                    route_in_str=""
                    route_out_str=""
                    route_name=""

                    if filtered_df is not None:
                        if (len(filtered_df)>0):
                            confidence_val="100"

                            if (len(filtered_df)>1):
                                print (f'Multiple Route found {row[uqfld] } between {frm_str_obj_id} AND {to_str_obj_id}')
                                logging.info(f'Multiple Route found {row[uqfld] } between {frm_str_obj_id} AND {to_str_obj_id}' )
                                selected_columns_df = filtered_df.loc[:, ['id','path','name','st_vntge_yr', 'created_at']]
                                vintage_yr = row['ST_VNTGE_YR']
                                create_dt= row['CREATED_DATE']
                                last_diff=999999999999999999999999
                                rwno=0
                                while rwno < len(selected_columns_df):
                                # for span_rw in filtered_df:
                                    if selected_columns_df.iloc[rwno,3]==vintage_yr:  #'st_vntge_yr'
                                        housing_route_id=selected_columns_df.iloc[rwno,0]#span_rw['id']
                                        geom_wkb=selected_columns_df.iloc[rwno,1]# span_rw['path']
                                        hoursing_route_name=selected_columns_df.iloc[rwno,2]
                                        route_in_str=filtered_df.iloc[rwno,2]
                                        route_out_str=filtered_df.iloc[rwno,3]
                                        break
                                    else:
                                        if selected_columns_df.iloc[rwno,4] ==create_dt:#span_rw['created_date']
                                            housing_route_id=selected_columns_df.iloc[rwno,0] #span_rw['id']
                                            geom_wkb= selected_columns_df.iloc[rwno,1] #span_rw['path']
                                            hoursing_route_name=selected_columns_df.iloc[rwno,2]
                                            route_in_str=filtered_df.iloc[rwno,2]
                                            route_out_str=filtered_df.iloc[rwno,3]

                                        else:
                                            if housing_route_id=="":
                                                #only applicable if housing id not found using created_date
                                                
                                                route_nm_last_part=util_services.UTILServices.get_char_after(selected_columns_df.iloc[rwno,2] ,':')#span_rw['name']
                                                route_nm_last_numpart= route_nm_last_part[1:]
                                                try:
                                                    diff  =  int(route_nm_last_numpart)-int(conduit_last_numpart)
                                                    if diff<0 : diff=diff*-1
                                                    if diff<last_diff:
                                                        last_diff=diff
                                                        housing_route_id=selected_columns_df.iloc[rwno,0] 
                                                        geom_wkb= selected_columns_df.iloc[rwno,1] 
                                                        hoursing_route_name=selected_columns_df.iloc[rwno,2]
                                                        route_in_str=filtered_df.iloc[rwno,2]
                                                        route_out_str=filtered_df.iloc[rwno,3]
                                                except Exception:
                                                    print('Not a numric value in span name.')
                                    rwno+=1

                                if housing_route_id=="":
                                    logging.info('No matching Vintage Year or Created Date found, using the 1st route as housing.' )
                                    print('No matching Vintage Year or Created Date found, using the 1st route as housing.')
                                    housing_route_id=filtered_df.iloc[0,0]
                                    geom_wkb=filtered_df.iloc[0,1]
                                    hoursing_route_name=filtered_df.iloc[0,4]
                                    route_in_str=filtered_df.iloc[0,2]
                                    route_out_str=filtered_df.iloc[0,3]

                            else:
                                # print ('Single Route found..')
                                housing_route_id=filtered_df.iloc[0,0]                                
                                geom_wkb=filtered_df.iloc[0,1]
                                hoursing_route_name=filtered_df.iloc[0,4]
                                route_in_str=filtered_df.iloc[0,2]
                                route_out_str=filtered_df.iloc[0,3]
                                
                        else:
                            confidence_val="0"
                            print('No route found for ' +row[uqfld] +' Adding route for conduit run process' )
                            logging.info('No route found for ' +row[uqfld] )
                            # logging.info( )
                            Missing_route_conduit.append([row[uqfld],frm_str_obj_id , to_str_obj_id,row['INVENTORY_STATUS_CODE'],row['ST_VNTGE_YR']])
                            # Missing_route_conduit.append([asso_span_row[17],frm_str_obj_id , to_str_obj_id])

                            # self.parent_missing.append([row[uqfld],"Housing / Root Housing",  frm_str_obj_id , to_str_obj_id])
                            continue
                    else:
                        continue
                 
                    #When housing route name is identified then find it's reference id used by IQGEO
                    route_match = df_route_ref[df_route_ref["name"].astype(str).str.strip().str.upper() == hoursing_route_name]
                    if not route_match.empty:
                        housing_route_id = str(route_match.iloc[0]["id"]).strip()
                        route_in_str=str(route_match.iloc[0]["in_structure"]).strip()
                        route_out_str=str(route_match.iloc[0]["out_structure"]).strip()
                    else:
                        confidence_val="0"
                        housing_route_id=hoursing_route_name
                        self.parent_missing.append([row[uqfld],"housing",hoursing_route_name, ""])
                    ########################
                    csv_row[1]=geom_wkb
                    csv_row[root_housing_idx]=housing_route_id
                    csv_row[housing_idx]=housing_route_id
                                        
                    for fld in self.filtered_fields:
                        if fld["Field_Name"].upper()=="FROM_STRUCTURE_NAME" : frm_str_fld_indx=fld_indx+2
                        if fld["Field_Name"].upper()=="TO_STRUCTURE_NAME": to_str_fld_indx=fld_indx+2
                        rw_fld_val=""
                        if fld["Field_Name"]=="name": rw_fld_val=asso_span_row[0]
                        if fld["Field_Name"]=="type_name": rw_fld_val=asso_span_row[1]
                        if fld["Field_Name"]=="specification": rw_fld_val=asso_span_row[2]
                        if fld["Field_Name"]=="inventory_status_code": rw_fld_val=asso_span_row[3]
                        if fld["Field_Name"]=="work_order_name": rw_fld_val=asso_span_row[4]
                        if fld["Field_Name"]=="st_std_desc": rw_fld_val=asso_span_row[5]
                        if fld["Field_Name"]=="measured_length": rw_fld_val=asso_span_row[6]
                        if fld["Field_Name"]=="ducts_available": rw_fld_val=asso_span_row[7]
                        if fld["Field_Name"]=="diameter": rw_fld_val=asso_span_row[8]
                        if fld["Field_Name"]=="label": rw_fld_val=asso_span_row[9]
                        if fld["Field_Name"]=="st_rmks": rw_fld_val=asso_span_row[10]
                        if fld["Field_Name"]=="st_vntge_yr": rw_fld_val=asso_span_row[11]
                        if fld["Field_Name"]=="sap_pm_order": rw_fld_val=asso_span_row[12]
                        if fld["Field_Name"]=="created_user": rw_fld_val=asso_span_row[13]
                        if fld["Field_Name"]=="created_at": rw_fld_val=asso_span_row[14]
                        if fld["Field_Name"]=="last_edited_date": rw_fld_val=asso_span_row[15]
                        if fld["Field_Name"]=="last_edited_date": rw_fld_val=asso_span_row[16]
                        if fld["Field_Name"]=="confidence_percentage": rw_fld_val=confidence_val

                        fld_val=self._ValidateFieldValue(fld,fld_indx,rw_fld_val,idval)
                        csv_row.append( fld_val)
                        fld_indx+=1

                    csv_row[frm_str_fld_indx]=route_in_str
                    csv_row[to_str_fld_indx]=route_out_str
        
                    self.records_skipped_due_to_issue=0
                    
                    if trunc_record==True :
                        self.number_of_records_truncated+=1
                    if db_type_issue==True :
                        self.records_skipped_due_to_issue+=1
                            
                    lengths = str_geom.length
                    v_n=  len(str_geom.coords)
                    FROM_STRUCTURE_NAME=row["FROM_STRUCTURE_NAME"]
                    TO_STRUCTURE_NAME=row["TO_STRUCTURE_NAME"]
                    if frm_str_obj_id>to_str_obj_id:
                        str_key=frm_str_obj_id+'|'+to_str_obj_id
                    else:
                        str_key=to_str_obj_id +'|'+frm_str_obj_id
                    
                        # cur.execute("INSERT INTO route_geom (str_hash , geom , lngth , no_vertex , obj_id , from_str_id , to_str_id , from_str_type , to_str_type , span_type ,in_str_name ,out_str_name  ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",(str_key,geom_wkb,lengths ,v_n ,iqgeo_ref_id_val, row['instr_object_id'],row['outstr_object_id'], row['INSTR_TYPE_NAME'],row['OUTSTR_TYPE_NAME'],row["TYPE_NAME"],FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME ))
                    cur.execute("INSERT INTO route_geom (str_hash , geom , lngth , no_vertex , obj_id , from_str_id , to_str_id , from_str_type , to_str_type , span_type ,in_str_name ,out_str_name ,is_active ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",(str_key,geom_wkb,lengths ,v_n ,iqgeo_ref_id_val, frm_str_obj_id,to_str_obj_id, row['INSTR_TYPE_NAME'],row['OUTSTR_TYPE_NAME'],row["TYPE_NAME"],FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME ,'Y'))
                    csv_data.append(csv_row)

            else:
                # idval=re.sub(r'[^0-9a-zA-Z]', '', row[uqfld])
                self.gen_conduit_id+=1
                iqgeo_ref_id_val= obj_id_pre_fix + '/'+ str(self.gen_conduit_id) 
                # if str(row[uqfld])=="{AE2DF7C2-6738-4035-8978-329F97B7E658}":
                idval=row["SPAN_NAME"]
                #     idval=re.sub(r'[^0-9a-zA-Z]', '', row[uqfld])

                # iqgeo_ref_id_val= obj_id_pre_fix +'/'+ str( idval)  
                # iqgeo_ref_id_val= obj_id_pre_fix + idval

                # uq_fld_val=row[src_uq_fld]
                str_geom=row["geometry"]  #geomfld
                str_geom=util_services.UTILServices.ConvertMultiLineToLine(str_geom)
                csv_row.append(iqgeo_ref_id_val)
                csv_row.append("")
                
                is_any_val_truncated=False
                fld_indx=0
                in_str_geom=None
                out_str_geom=None
                for fld in self.reference_fields:
                    fld_val =""
                    if fld['Field_Name']=="in_structure":
                        str_name=row[self.in_structure_left_clm.upper()]
                        match = df_str_ref[df_str_ref["name"].astype(str).str.strip().str.upper() == str_name]
                        if not match.empty:
                            structure_id = str(match.iloc[0]["id"]).strip()
                            fld_val = structure_id
                            frm_str_obj_id = structure_id
                            # in_str_geom=row["InStrGeom"]
                        else:
                            fld_val = f"{str_name}"#str_name
                            # fld_val=str_name
                            frm_str_obj_id = ""
                            self.parent_missing.append([row[uqfld],"housing",row[self.in_structure_right_clm + "_x"],row[self.in_structure_right_clm + "_y"]])

                    if fld['Field_Name']=="out_structure":
                        str_name=row[self.out_structure_left_clm.upper()]
                        match = df_str_ref[df_str_ref["name"].astype(str).str.strip().str.upper() == str_name]
                        if not match.empty:
                            structure_id = str(match.iloc[0]["id"]).strip()
                            fld_val = structure_id
                            to_str_obj_id = structure_id
                            # out_str_geom=row["OutStrGeom"]
                        else:
                            fld_val = f"{str_name}"#str_name
                            # fld_val=str_name
                            to_str_obj_id = ""
                            self.parent_missing.append([row[uqfld],"housing",row[self.out_structure_right_clm + "_x"],row[self.out_structure_right_clm + "_y"]])


                    # if fld['Field_Name']=="in_structure":
                    #     if pd.isnull(row["instr_object_id"]):
                    #         frm_str_obj_id="" 
                    #         self.parent_missing.append([row[uqfld],"in_structure", row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                    #     else:
                    #         frm_str_obj_id=re.sub(r'[^0-9a-zA-Z]', '', row["instr_object_id"])  

                    #         in_str_geom=row["InStrGeom"]
                    #         validated_obj=util_services.UTILServices.check_and_convert_value(fld_val,"str")
                    #         if validated_obj['is_null']==True or validated_obj['is_nan']  ==True:
                    #             frm_str_obj_id=""
                    #             self.parent_missing.append([row[uqfld],"in_structure", row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                    #         else:
                    #             frm_str_obj_id =self.in_structure_fc.upper()+str((frm_str_obj_id))
                    #     fld_val=frm_str_obj_id

                    # if fld['Field_Name']=="out_structure":
                    #     if pd.isnull(row["outstr_object_id"]):
                    #         to_str_obj_id="" 
                    #         self.parent_missing.append([row[uqfld],"out_structure",  row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                    #     else:
                    #         to_str_obj_id=re.sub(r'[^0-9a-zA-Z]', '', row["outstr_object_id"])   
                    #         out_str_geom=row["OutStrGeom"]

                    #         validated_obj=util_services.UTILServices.check_and_convert_value(fld_val,"str")
                    #         if validated_obj['is_null']==True or validated_obj['is_nan']  ==True:
                    #             to_str_obj_id=""
                    #             self.parent_missing.append([row[uqfld],"out_structure",  row[self.in_structure_right_clm+'_x'] , row[self.out_structure_right_clm +'_x']])
                    #         else:
                    #             to_str_obj_id =self.in_structure_fc.upper()+str((to_str_obj_id))
                    #     fld_val=to_str_obj_id

                    csv_row.append( fld_val)
                    fld_indx+=1

    # Read Route CSV for Housing reference
                conduit_last_part=util_services.UTILServices.get_char_after(row['SPAN_NAME'],':')
                conduit_last_numpart= conduit_last_part[1:]
                combined_conditions  = f"((route_df['in_structure'] =='{frm_str_obj_id}') &  (route_df['out_structure'] ==  '{to_str_obj_id }'))  | ((route_df['in_structure'] =='{to_str_obj_id}') & (route_df['out_structure'] ==  '{frm_str_obj_id}'))"
                filtered_df = route_df[eval(combined_conditions)]
                route_in_str=""
                route_out_str=""
                if filtered_df is not None:
                    if (len(filtered_df)>0):
                        confidence_val="100"

                        if (len(filtered_df)>1):
                            print (f'Multiple Route found {row[uqfld] } between {frm_str_obj_id} AND {to_str_obj_id}')
                            logging.info(f'Multiple Route found {row[uqfld] } between {frm_str_obj_id} AND {to_str_obj_id}' )
                            selected_columns_df = filtered_df.loc[:, ['id','path','name','st_vntge_yr', 'created_at']]
                            vintage_yr = row['ST_VNTGE_YR']
                            create_dt= row['CREATED_DATE']
                            span_nm= row['SPAN_NAME']
                            last_diff=999999999999999999999999
                            rwno=0
                            while rwno < len(selected_columns_df):
                            # for span_rw in filtered_df:
                                if selected_columns_df.iloc[rwno,3]==vintage_yr:  #'st_vntge_yr'
                                    housing_route_id=selected_columns_df.iloc[rwno,0]#span_rw['id']
                                    geom_wkb=selected_columns_df.iloc[rwno,1]# span_rw['path']
                                    hoursing_route_name=selected_columns_df.iloc[rwno,2]
                                    route_in_str=filtered_df.iloc[rwno,2]
                                    route_out_str=filtered_df.iloc[rwno,3]
                                    break
                                else:
                                    if selected_columns_df.iloc[rwno,4] ==create_dt:#span_rw['created_date']
                                        housing_route_id=selected_columns_df.iloc[rwno,0] #span_rw['id']
                                        geom_wkb= selected_columns_df.iloc[rwno,1] #span_rw['path']
                                        hoursing_route_name=selected_columns_df.iloc[rwno,2]
                                        route_in_str=filtered_df.iloc[rwno,2]
                                        route_out_str=filtered_df.iloc[rwno,3]
                                    else:
                                        if housing_route_id=="":
                                            #only applicable if housing id not found using created_date
                                            
                                            route_nm_last_part=util_services.UTILServices.get_char_after(selected_columns_df.iloc[rwno,2] ,':')#span_rw['name']
                                            route_nm_last_numpart= route_nm_last_part[1:]
                                            try:
                                                diff  =  int(route_nm_last_numpart)-int(conduit_last_numpart)
                                                if diff<0 : diff=diff*-1
                                                if diff<last_diff:
                                                    last_diff=diff
                                                    housing_route_id=selected_columns_df.iloc[rwno,0] 
                                                    geom_wkb= selected_columns_df.iloc[rwno,1] 
                                                    hoursing_route_name=selected_columns_df.iloc[rwno,2]
                                                    route_in_str=filtered_df.iloc[rwno,2]
                                                    route_out_str=filtered_df.iloc[rwno,3]
                                            except Exception:
                                                print('Not a numric value in span name.')
                                rwno+=1

                            if housing_route_id=="":
                                logging.info('No matching Vintage Year or Created Date found, using the 1st route as housing.' )
                                print('No matching Vintage Year or Created Date found, using the 1st route as housing.')
                                housing_route_id=filtered_df.iloc[0,0]
                                geom_wkb=filtered_df.iloc[0,1]
                                hoursing_route_name=filtered_df.iloc[0,4]                                
                                route_in_str=filtered_df.iloc[0,2]
                                route_out_str=filtered_df.iloc[0,3]
                        else:
                            # print ('Single Route found..')
                            housing_route_id=filtered_df.iloc[0,0]
                            geom_wkb=filtered_df.iloc[0,1]
                            hoursing_route_name=filtered_df.iloc[0,4]
                            route_in_str=filtered_df.iloc[0,2]
                            route_out_str=filtered_df.iloc[0,3]
                    else:
                        confidence_val="0"
                        print('No route found for ' +row[uqfld] +' Adding route for conduit run process' )
                        logging.info('No route found for ' +row[uqfld] )
                        Missing_route_conduit.append([row[uqfld],frm_str_obj_id , to_str_obj_id,row['INVENTORY_STATUS_CODE'],row['ST_VNTGE_YR']])

                        # logging.info( )

                        # self.parent_missing.append([row[uqfld],"Housing / Root Housing",  frm_str_obj_id , to_str_obj_id])
                        continue
                else:
                    continue
                # if in_str_geom !=None:
                #     if out_str_geom !=None:
                #         in_str_type = row['TYPE_NAME']
                #         str_geom,d1,d2= util_services.UTILServices.change_path_start_endpoints(str_geom,in_str_geom, out_str_geom)
                #         if d1>10 or d2>10:
                #             self.topology_issues.append([idval,d1,d2]) 
                
                # geom_wkb=util_services.UTILServices.getEWKB(str_geom,src_crs,tgt_crs)

                route_match = df_route_ref[df_route_ref["name"].astype(str).str.strip().str.upper() == hoursing_route_name]
                if not route_match.empty:
                    housing_route_id = str(route_match.iloc[0]["id"]).strip()
                    route_in_str=str(route_match.iloc[0]["in_structure"]).strip()
                    route_out_str=str(route_match.iloc[0]["out_structure"]).strip()
                else:
                    housing_route_id=hoursing_route_name
                    self.parent_missing.append([row[uqfld],"housing",hoursing_route_name, ""])

                csv_row[1]=geom_wkb
                csv_row[root_housing_idx]=housing_route_id
                csv_row[housing_idx]=housing_route_id
                
                
            # fld_indx+=1

                for fld in self.filtered_fields:
                    if fld["Field_Name"].upper()=="FROM_STRUCTURE_NAME" : frm_str_fld_indx=fld_indx+2
                    if fld["Field_Name"].upper()=="TO_STRUCTURE_NAME": to_str_fld_indx=fld_indx+2

                    fld_val=self._GetFieldValue(fld,fld_indx,row,idval)
                    csv_row.append( fld_val)
                    fld_indx+=1
    
                self.records_skipped_due_to_issue=0
                #setting up in /out structure as per the parent route not as per the NE data
                csv_row[frm_str_fld_indx]=route_in_str
                csv_row[to_str_fld_indx]=route_out_str
                csv_row[conf_val_idx]=confidence_val

                ######################
                if trunc_record==True :
                    self.number_of_records_truncated+=1
                if db_type_issue==True :
                    self.records_skipped_due_to_issue+=1
                        
                lengths = str_geom.length
                v_n=  len(str_geom.coords)
                FROM_STRUCTURE_NAME=row["FROM_STRUCTURE_NAME"]
                TO_STRUCTURE_NAME=row["TO_STRUCTURE_NAME"]
                if frm_str_obj_id>to_str_obj_id:
                    str_key=frm_str_obj_id+'|'+to_str_obj_id
                else:
                    str_key=to_str_obj_id +'|'+frm_str_obj_id
                csv_data.append(csv_row)
                    # cur.execute("INSERT INTO route_geom (str_hash , geom , lngth , no_vertex , obj_id , from_str_id , to_str_id , from_str_type , to_str_type , span_type ,in_str_name ,out_str_name  ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",(str_key,geom_wkb,lengths ,v_n ,iqgeo_ref_id_val, row['instr_object_id'],row['outstr_object_id'], row['INSTR_TYPE_NAME'],row['OUTSTR_TYPE_NAME'],row["TYPE_NAME"],FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME ))
                cur.execute("INSERT INTO route_geom (str_hash , geom , lngth , no_vertex , obj_id , from_str_id , to_str_id , from_str_type , to_str_type , span_type ,in_str_name ,out_str_name ,is_active ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",(str_key,geom_wkb,lengths ,v_n ,iqgeo_ref_id_val, frm_str_obj_id,to_str_obj_id, row['INSTR_TYPE_NAME'],row['OUTSTR_TYPE_NAME'],row["TYPE_NAME"],FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME ,'Y'))
        

            rec_processed+=1
            percent =int(100* (rec_processed / total_records))
            print (f"{rec_processed} of {total_records} - {percent} %",end="\r")
        
        if self.category=="CONDUIT":
            formation_rw=self.GetFormationSpanForConduit(route_df,cur,csv_data,obj_id_pre_fix,df_route_ref)
            rec_processed= len(formation_rw)
            
         
        conn.commit()
        conn2.close()
# Checking Messenger data and modifying geom to shortest path if found
        span_name_index=util_services.UTILServices.GetColumnIndex(clm_name, "name")
        # if len(csv_data)>0:
        #     conduit_conduitrun_data,cr_data= self._CreateConduitRun(csv_data,clm_name,span_name_index)
        # max_id_conduit_run = len(cr_data)
        clidx=0
        

        nm_idx=-1
        frm_str_type_id = -1
        to_str_type_id = -1
        in_counduit_idx =util_services.UTILServices.GetColumnIndex(clm_name,'in_conduit')
        out_counduit_idx= util_services.UTILServices.GetColumnIndex(clm_name,'out_conduit')
        if frm_str_fld_indx== -1 or to_str_fld_indx==-1 :
            for cl in clm_name:
                if cl=='in_structure':frm_str_fld_indx= clidx
                if cl=='out_structure':to_str_fld_indx= clidx
                if cl=='name':
                    nm_idx=clidx
                clidx+=1
        if len(Missing_route_conduit)>0:
            Missing_route_conduit=util_services.UTILServices.remove_duplicates(Missing_route_conduit)
            for Missing_conduit in Missing_route_conduit:
                # if Missing_conduit[2]=="STRUCTURE635693":
                #     Missing_conduit[2]="STRUCTURE256857"
                # if Missing_conduit[1]=="STRUCTURE635706":
                #     Missing_conduit[1]="STRUCTURE250481"
                confidence_val="0"
                conduit_uuid = Missing_conduit[0]
                print(f"Missing path for conduit  uuid->{conduit_uuid} between {Missing_conduit[1]} and {Missing_conduit[2]}")
                # Filter for conduit uuid to populate fields value
                conduit_gdf = self.gdf[self.gdf["object_id"]== conduit_uuid]
                if len(conduit_gdf)>0:
                    row=conduit_gdf.iloc[0] #get refrence of 1st row of filtered data frame
                    if row['SPAN_NAME']=='OSP:COND::12730190':
                        print('Break for debug')
                    length_of_org_conduit= row["geometry"].length
                    
                    route_df['year'] = route_df['st_vntge_yr'].fillna(0).astype(int)

                    G = self.set_G_edges(route_df)  
                    data = self.get_shortest_path(route_df, Missing_conduit[1], Missing_conduit[2],G,Missing_conduit[3],Missing_conduit[4])
                    confidence_val = data.get('confidence_percentage')
                    if "error" in data.keys():
                        # data = self.get_shortest_path(route_df, to_str_obj_id, frm_str_obj_id)
                        data = self.get_shortest_path(route_df, Missing_conduit[2], Missing_conduit[1],G,Missing_conduit[3],Missing_conduit[4])
                        confidence_val = data.get('confidence_percentage')

                    if "error" in data.keys():
                        self.parent_missing.append([row['SPAN_NAME'],"Unable to get route for conduit",  Missing_conduit[1] , Missing_conduit[2]])
                    else:
                        # print(f"Found shortest path, processing data {data}")
                        #geomerty
                        geom_list = []
                        length_of_all_parts=0
                        for route_name in data['path_name']:

                            route_rw = route_df.loc[route_df['name'] == route_name].iloc[0]
                            route_geom = route_rw["path"] #route_df.loc[route_df['name'] == route_name, 'path'].iloc[0]
                            
                            geom_list.append(route_geom)            
                        
                        try:
                            merged_geom = util_services.UTILServices.merge_geom(geom_list)
                        except ValueError:
                            print("Warning: Cannot linemerge disconnected lines, using MultiLineString fallback.")
                            
                            
                            geom_objects = [wkb.loads(bytes.fromhex(g)) for g in geom_list]
                            merged_geom_obj = MultiLineString(geom_objects)
                            merged_geom = merged_geom_obj.wkb.hex().upper()

                        merged_geom_object = wkb.loads(merged_geom,hex=True)

                        length_of_all_parts=merged_geom_object.length
                        
                        if length_of_all_parts<length_of_org_conduit*1.3:
                            max_id_conduit_run +=1
                            cr_id='mywcom_conduit_run/'+str(max_id_conduit_run)
                            additional_conduit_run.append([cr_id,merged_geom])
                            # conduit_rw[conduitrun_col_index]=cr_id
                            # csv_row[1] = merged_geom
                            i=0
                            for route_name in data['path_name']:
                                self.gen_conduit_id+=1
                                i+=1
                                route_rw = route_df.loc[route_df['name'] == route_name].iloc[0]
                                route_geom = route_rw["path"] #route_df.loc[route_df['name'] == route_name, 'path'].iloc[0]

                                route_match = df_route_ref[df_route_ref["name"].astype(str).str.strip().str.upper() == route_name]
                                if not route_match.empty:
                                    housing_route_id = str(route_match.iloc[0]["id"]).strip()
                                    route_in_str=str(route_match.iloc[0]["in_structure"]).strip()
                                    route_out_str=str(route_match.iloc[0]["out_structure"]).strip()
                                    # confidence_val="90"
                                else:
                                    housing_route_id=hoursing_route_name
                                    self.parent_missing.append([row[uqfld],"housing",hoursing_route_name, ""])
                                    route_in_str=""
                                    route_out_str="" 
                                    

                                geom_list.append(route_geom)
                                csv_row=[]
                                idval= route_rw[0]
                                iqgeo_ref_id_val= obj_id_pre_fix + '/' +str( self.gen_conduit_id  )


                                csv_row.append(iqgeo_ref_id_val)
                                csv_row.append(route_geom)

                                # csv_row.append(route_rw['id'])  # housing
                                # csv_row.append(route_rw['id'])  # root housing
                                # csv_row.append(route_rw['in_structure'])
                                # csv_row.append(route_rw['out_structure'])
                                csv_row.append(housing_route_id)  # housing
                                csv_row.append(housing_route_id)  # root housing
                                csv_row.append(route_in_str)
                                csv_row.append(route_out_str)
                                
                                csv_row.append("")  # in_conduit
                                csv_row.append("")  # out_conduit
                                csv_row.append(cr_id)  # conduit_run
                                
                                fld_indx=7
                                conduit_name=""
                                for fld in self.filtered_fields:
                                    
                                    fld_val=self._GetFieldValue(fld,fld_indx,row,idval)
                                    if fld["Field_Name"]=="name": 
                                        conduit_name=fld_val+'-'+str(i)
                                        fld_val=conduit_name

                                    if fld["Field_Name"]=="confidence_percentage": fld_val=confidence_val

                                    csv_row.append( fld_val)
                                    fld_indx+=1

                                #checking conduit name is already added in csv and if YES then just bypass
                                searched_item = [r for r in csv_data if r[9]==conduit_name]
                                if searched_item:
                                    print ('bypass')
                                else:
                                    csv_data.append(csv_row)
                                cur.execute("INSERT INTO route_geom (  obj_id , from_str_id , to_str_id ,cr_id  ) VALUES (?,?,?,? )",( iqgeo_ref_id_val, route_rw['in_structure'],route_rw['out_structure'],cr_id))
                        else:
                            #Length of identified shortest route is more than 30% of original conduit route hence not considered as correct route
                            
                            self.parent_missing.append([row['SPAN_NAME'],"Shortest route not found", 'Identified route is 30 percent longer then orginal counduit route ' ,data['path_name'] ])
                    
                #Identify routes
                #getconduit for each route
                #create conduit run by merging the geomwtry of the routes
                #define in conduit an d out conduit
                #add into thr conduit_rw




        print("Processing for In /Out Counduit referencing...Please wait it will take some time..")
        logging.info("Processing for In /Out Counduit referencing..." )
        # conduitrun_df = pd.DataFrame(conduit_conduitrun_data)
        conduitrun_col_index=util_services.UTILServices.GetColumnIndex(clm_name, "conduit_run")

        for conduit_rw in csv_data:
            FROM_STRUCTURE_NAME=conduit_rw[frm_str_fld_indx]
            TO_STRUCTURE_NAME=conduit_rw[to_str_fld_indx]
            id_val=conduit_rw[0]
            cr_id = conduit_rw[conduitrun_col_index]
            # span_name = conduit_rw[span_name_index]
            # if len(conduitrun_df)==0 or 2 not in conduitrun_df.columns:
            #     print("skipping in out conduit")
            #     continue
            # filtered_df = conduitrun_df[conduitrun_df[2] ==span_name]
            # if len(filtered_df)>0:
                # conduit_rw[conduitrun_col_index]=filtered_df.iloc[0,0]
            query = f"SELECT  obj_id FROM  route_geom WHERE  to_str_id = '{FROM_STRUCTURE_NAME}' and cr_id = '{cr_id}' "
            cur.execute(query )
            results = cur.fetchall()
            if results:
                if len(results)>1:
                    #multiple in_counduit available, identify which one to be used. 
                    conduit_rw[in_counduit_idx]=results[0][0]
                    print(f"Multiple in_conduit found for {id_val}")
                else:
                    conduit_rw[in_counduit_idx]=results[0][0]
            else:
                print(f"No In Counduit found for {id_val}")
                logging.info(f"No In Counduit found for {id_val}" )

            query = f"SELECT  obj_id FROM  route_geom WHERE  from_str_id = '{TO_STRUCTURE_NAME}'  and cr_id = '{cr_id}'"
        
            cur.execute(query )
            results = cur.fetchall()
            if results:
                if len(results)>1:
                    #multiple out_counduit available, identify which one to be used. 
                    print(f"Multiple out_counduit found for {id_val}")

                    conduit_rw[out_counduit_idx]=results[0][0]
                else:
                    conduit_rw[out_counduit_idx]=results[0][0]

            else:
                print(f"No out Counduit found for {id_val}")
                logging.info(f"No Out Counduit found for {id_val}" )    
            # else:
            #     ## In case of no no conduit_run is generated then adding a default conduit_run  
            #         ## Reason to add- Conduit Run is mandatory in IQGeo model and due to this showing conflict 
            #         ## hence it is decided to add same geometry of conduit in its conduit_run
            #         ## 23 Oct 2025
            #         max_id_conduit_run +=1
            #         cr_id='mywcom_conduit_run/'+str(max_id_conduit_run)
            #         additional_conduit_run.append([cr_id,merged_geom])
            #         conduit_rw[conduitrun_col_index]=cr_id
            if conduit_rw[conduitrun_col_index]=="":
                #=====================================
                # Adding same path in conduit_run for all such conduits those do not have conduit_run as per the above processes
                max_id_conduit_run +=1
                cr_id='mywcom_conduit_run/'+str(max_id_conduit_run)
                conduit_rw[conduitrun_col_index]=cr_id
                additional_conduit_run.append([cr_id,conduit_rw[1]])

        if len(additional_conduit_run)>0:
                add_df=pd.DataFrame(additional_conduit_run,columns=['id','path'])
                # cr_df=pd.DataFrame(cr_data,columns=['id','path'])

                # merged_df = pd.concat([cr_df, add_df],ignore_index=True)
                conduit_run_file_csv=self.folder_path+"\\mywcom_conduit_run.csv"
                add_df.to_csv(conduit_run_file_csv,index=False)
                print(f'Writing conduit_run csv file.{conduit_run_file_csv}')
                #self._WriteCSVFile(conduit_run_file_csv ,['id','path'],cr_data)   
                try:
                    rel = os.path.relpath(conduit_run_file_csv, self.rootDirectory)
                    self._created_files.append(rel)
                except Exception:
                    pass     
#=====================================
        conn.commit()
        conn.close()
        try:
            os.remove(sqlite_path)
            print('Removing db file')
        except Exception as ex:
            print(f"Unable to delete temp db  file {sqlite_path}")
            
        return csv_data,rec_processed
    def GetFormationSpanForConduit(self,route_df,cur,formation_csv,obj_id_pre_fix,df_route_ref):
        # formation_csv=[]
        print("Processing for Formation type ...Please wait it will take some time..")
        logging.info("Processing for Formation Type..." )

        clm_names= self.NE_Fields
        # params = oracledb.ConnectParams(host="sod-oradb-023.stholdco.com", port=1527, service_name="DBD463.STHOLDCO.COM")
        combined_conditions  = f"(route_df['type_name'] =='FORMATION')  "
        filtered_df = route_df[eval(combined_conditions)]
        filtered_cnt=len(filtered_df)
        print(f'Total {filtered_cnt} Formation type found.')
        if filtered_df is not None:
            if (filtered_cnt>0):
                rwno=0
                comma_separated_clms = ",su.".join(clm_names)
                sql_conn = sqlite3.connect(self.sqlite_db)
                sql_cursor = sql_conn.cursor()
                # ora_conn = oracledb.connect(user="xxxxxxxxxxx", password="xxxxxxxx", params=params)
                # ora_cursor = ora_conn.cursor()
                for index, row in filtered_df.iterrows():
                    rwno+=1
                    # formation_id = row['id']
                    print (f"Processing {rwno} of {filtered_cnt}",end="\r")
                    confidence_val="0"
                    formation_span_name= row['name']
                    # if formation_span_name =='OSP:FORSP::13702967' or formation_span_name=='OSP:FORSP::13702966':
                    #     print('reached')
                    formation_geom=util_services.UTILServices._load_wkb( row['path'])
                    formation_geom_wkb= row['path']

                    route_match = df_route_ref[df_route_ref["name"].astype(str).str.strip().str.upper() == formation_span_name]
                    if not route_match.empty:
                        housing_route_id = str(route_match.iloc[0]["id"]).strip()
                        route_in_str=str(route_match.iloc[0]["in_structure"]).strip()
                        route_out_str=str(route_match.iloc[0]["out_structure"]).strip()
                        confidence_val="100"

                    else:
                        housing_route_id=formation_span_name
                        self.parent_missing.append([formation_span_name,"housing",formation_span_name, ""])
                        route_in_str=""
                        route_out_str=""

                    # in_str_id = row['in_structure']
                    # out_str_id = row['out_structure']
                    # if pd.isnull(in_str_id): in_str_id=""
                    # if pd.isnull(out_str_id): out_str_id=""
                    
                    #sql_query = f"SELECT su.uuid, {comma_separated_clms} FROM ne.SPAN_SPAN_UNIT ssu INNER JOIN ne.SPAN_UNIT su  ON ssu.SPAN_UNIT_UUID =su.UUID  INNER JOIN ne.SPAN s  ON s.UUID = ssu.SPAN_UUID WHERE s.SPAN_NAME ='{formation_span_name}'"
                    # spanunit_df = pd.read_sql(sql_query, conn)
                    # sql_query = f"SELECT DISTINCT su.uuid, su.type_name, s.span_ref_name, su.inventory_status_code,su.work_order_name,su.account_code,su.measured_length,su.ducts_available,su.diameter,su.label,su.st_rmks,su.st_vntge_yr,su.st_sap_ntwk_id,su.FROM_STRUCTURE_NAME,su.TO_STRUCTURE_NAME , su.SPAN_UNIT_NAME SPAN_UNIT FROM ne.MV_SPAN_SPAN_UNIT ssu INNER JOIN ne.MV_SPAN_UNIT su  ON ssu.SPAN_UNIT_UUID =su.UUID  INNER JOIN ne.MV_SPAN s  ON s.UUID = ssu.SPAN_UUID WHERE su.TYPE_NAME<>'CORE' AND  s.SPAN_NAME ='{formation_span_name}'" sql_cursor.close()
                    sql_query = f"SELECT DISTINCT su.uuid, su.type_name, s.span_ref_name, su.inventory_status_code,su.work_order_name,su.account_code,su.measured_length,su.ducts_available,su.diameter,su.label,su.st_rmks,su.st_vntge_yr,su.st_sap_ntwk_id,su.FROM_STRUCTURE_NAME,su.TO_STRUCTURE_NAME , su.SPAN_UNIT_NAME SPAN_UNIT FROM SPAN_SPAN_UNIT ssu INNER JOIN SPAN_UNIT su  ON ssu.SPAN_UNIT_UUID =su.UUID  INNER JOIN SPAN s  ON s.UUID = ssu.SPAN_UUID WHERE su.TYPE_NAME<>'CORE' AND  s.SPAN_NAME ='{formation_span_name}'" 
                    # sql_cursor.close()
        
                    # ora_cursor.execute(sql_query)
                    # rows = ora_cursor.fetchall()
                    sql_cursor.execute(sql_query)
                    rows = sql_cursor.fetchall()
                    for span_unit_row in rows:
                        fld_indx=0
                        csv_row=[]
                        self.gen_conduit_id+=1
                        conduit_id=obj_id_pre_fix +  '/'+str(self.gen_conduit_id)  #'uuid' 
                        # conduit_id=obj_id_pre_fix + re.sub(r'[^0-9a-zA-Z]', '', span_unit_row[0])  #'uuid' 
                        csv_row.append(conduit_id)
                        csv_row.append(formation_geom_wkb)


                        for fld in self.reference_fields:
                            fld_val =""
                            if fld['Field_Name']=="in_structure": fld_val=route_in_str
                            if fld['Field_Name']=="out_structure": fld_val=route_out_str
                            if fld['Field_Name']=="housing": fld_val=housing_route_id
                            if fld['Field_Name']=="root_housing": fld_val=housing_route_id

                            csv_row.append( fld_val)
                            fld_indx+=1
 
                        for fld in self.filtered_fields:
                            fld_val=""
                            if fld['Field_Name']=="type_name":                  rw={"TYPE_NAME":span_unit_row[1]}    #fld_val=span_unit_row[1]
                            if fld['Field_Name']=="span_ref_name":              rw={"SPAN_REF_NAME":span_unit_row[2]}                     #fld_val=span_unit_row[2]
                            if fld['Field_Name']=="inventory_status_code":      rw={"INVENTORY_STATUS_CODE":span_unit_row[3]}                     #fld_val=span_unit_row[3]
                            if fld['Field_Name']=="work_order_name":            rw={"WORK_ORDER_NAME":span_unit_row[4]}                     #fld_val=span_unit_row[4]
                            if fld['Field_Name']=="st_std_desc":               rw={"ACCOUNT_CODE":span_unit_row[5]}                     #fld_val=span_unit_row[5]
                            if fld['Field_Name']=="measured_length":            rw={"MEASURED_LENGTH":span_unit_row[6]}                     #fld_val=span_unit_row[6]
                            if fld['Field_Name']=="ducts_available":            rw={"DUCTS_AVAILABLE":span_unit_row[7]}                     #fld_val=span_unit_row[7]
                            if fld['Field_Name']=="diameter":                   rw={"DIAMETER":span_unit_row[8]}               #fld_val=span_unit_row[8]
                            if fld['Field_Name']=="label":                      rw={"LABEL":span_unit_row[9]}            #fld_val=span_unit_row[9]
                            if fld['Field_Name']=="st_rmks":                    rw={"ST_RMKS":span_unit_row[10]}             # fld_val=span_unit_row[10]
                            if fld['Field_Name']=="st_vntge_yr":     
                                rw={"ST_VNTGE_YR":span_unit_row[11]}              #    fld_val=span_unit_row[11]
                            if fld['Field_Name']=="sap_pm_order":        
                                rw={"ST_SAP_NTWK_ID":span_unit_row[12]}               #      fld_val=span_unit_row[12]
                            
                            if fld['Field_Name']=="name":            
                                rw={"SPAN_NAME":span_unit_row[15]}               #      fld_val=span_unit_row[12]
# account_code  st_std_desc
                            if fld['Field_Name']=="confidence_percentage":            
                                rw={"CONFIDENCE_PERCENTAGE":confidence_val}  

                            fld_val=self._GetFieldValue(fld,fld_indx,rw,conduit_id)
                            # fld_val= span_unit_row[fld]
                            csv_row.append( fld_val)
                            fld_indx+=1
                        lengths = formation_geom.length
                        v_n=  len(formation_geom.coords)
                        FROM_STRUCTURE_NAME=span_unit_row[13]
                        TO_STRUCTURE_NAME=span_unit_row[14]
                        if route_in_str>route_out_str:
                            str_key=route_in_str+'|'+route_out_str
                        else:
                            str_key=route_out_str +'|'+route_in_str
                        formation_csv.append(csv_row)
                            # cur.execute("INSERT INTO route_geom (str_hash , geom , lngth , no_vertex , obj_id , from_str_id , to_str_id , from_str_type , to_str_type , span_type ,in_str_name ,out_str_name  ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",(str_key,geom_wkb,lengths ,v_n ,iqgeo_ref_id_val, row['instr_object_id'],row['outstr_object_id'], row['INSTR_TYPE_NAME'],row['OUTSTR_TYPE_NAME'],row["TYPE_NAME"],FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME ))
                        cur.execute("INSERT INTO route_geom (str_hash , geom , lngth , no_vertex , obj_id , from_str_id , to_str_id , from_str_type , to_str_type , span_type ,in_str_name ,out_str_name ,is_active ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",(str_key,formation_geom_wkb,lengths ,v_n ,conduit_id, route_in_str,route_out_str, 'INSTR_TYPE_NAME','OUTSTR_TYPE_NAME','TYPE_NAME',FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME ,'Y'))
                # ora_cursor.close()
                # ora_conn.close()
                sql_cursor.close()
                sql_conn.close()
        
        return  formation_csv
            #(DESCRIPTION=(ADDRESS_LIST=(ADDRESS=(PROTOCOL=TCP)(HOST=sod-oradb-023.stholdco.com)(PORT=1527)))(CONNECT_DATA=(SERVICE_NAME=DBD463.STHOLDCO.COM)))


        #     # self.org_gdf=gpd.read_file(self.input_gdb,layer=self.src_lyr,columns=NE_Fields_Name)

        #     print(f"GDB read completed at {time.time()}")
            
        #     print("Filtering for  " + self.filter_column) 
        #     mergedgdf=util_services.UTILServices.apply_sql_like_filter(self.org_gdf,self.filter_column) 

        #     if in_structure_fc!="": #In structure joining needs to be done
        #         instr_gdf=gpd.read_file(self.input_gdb,layer=right_table,  columns=['object_id', in_structure_right_clm,'TYPE_NAME'])
        #         instr_gdf= instr_gdf.rename(columns={'object_id': 'instr_object_id', in_structure_right_clm: 'INSTR_'+in_structure_right_clm,'TYPE_NAME': 'INSTR_TYPE_NAME'})   

        #         self.in_structure_right_clm='INSTR_'+in_structure_right_clm

        #         instr_gdf_atr=instr_gdf.drop(columns='geometry')
        #         instr_gdf_geom=instr_gdf[[self.in_structure_right_clm, 'geometry']].rename(columns={"geometry":"InStrGeom"}) 
        #         mergedgdf=mergedgdf.merge(instr_gdf_atr,  left_on=in_structure_left_clm, right_on=self.in_structure_right_clm, how='left' ).merge(instr_gdf_geom, left_on=in_structure_left_clm, right_on=self.in_structure_right_clm, how='left')
                



      

    # def set_G_edges(self,df):
    #     G = nx.Graph()
    #     try:
    #         for _, row in df.iterrows():G.add_edge(row['in_structure'], row['out_structure'],weight = wkb.  loads(bytes.fromhex(row['path'])).length, id = row['id'])
    #     except Exception as e:
    #         return {"error": str(e)}
    
    #     return G
    def set_G_edges(self,df, filter=False, filter_by_path_name=None, vintage_yr = None):
        G = nx.Graph()
        try:
            if filter:
                for _, row in df.iterrows():
                    if row['name'] in filter_by_path_name:
                        if vintage_yr is None or row['year'] == vintage_yr:
                            G.add_edge(row['in_structure'], row['out_structure'],weight = wkb.  loads(bytes.fromhex(row['path'])).length, id = row['id'])
            else:
                for _, row in df.iterrows():G.add_edge(row['in_structure'], row['out_structure'],weight = wkb.  loads(bytes.fromhex(row['path'])).length, id = row['id'])
        except Exception as e:
            return {"error": str(e)}
        return G
       

    def get_shortest_path_edges(self,G, df, source_node, destination_node, ivs, vintage_yr):
        # return nx.shortest_path(G, source=source_node, target=destination_node, weight='weight')
        paths_generator = nx.all_simple_paths(G,source=source_node,target=destination_node,cutoff=8)
        all_path = list(paths_generator)
        if len(all_path)==0:
            print('No path segment found upto 8, now checking with 15')
            paths_generator = nx.all_simple_paths(G,source=source_node,target=destination_node,cutoff=15)
            all_path = list(paths_generator)
            if len(all_path)==0:
                print('No path segment found upto 15, now checking with 25')
                paths_generator = nx.all_simple_paths(G,source=source_node,target=destination_node,cutoff=25)
                all_path = list(paths_generator)

        # if len(all_path) == 1:
        #     return all_path[0]
        filter_by_path_name = []
        path_list = []
        for count, i in enumerate(all_path):
            check_status_code = []
            check_st_vntge_yr = []
            for j in range(0, len(i)-1):
                check_status_code.append(self.get_path_(df, [i[j], i[j+1]], 'inventory_status_code'))
                check_st_vntge_yr.append(0 if np.isnan(self.get_path_(df, [i[j], i[j+1]], 'st_vntge_yr')) else int(self.get_path_(df, [i[j], i[j+1]], 'st_vntge_yr')))
                # check_st_vntge_yr.append(get_path_(df, [i[j], i[j+1]], 'year'))
            if len(set(check_status_code)) == 1:
                if ivs == 'PLA':
                    if ivs == next(iter(set(check_status_code))) or 'PDA' == next(iter(set(check_status_code))) \
                        or 'IPL' == next(iter(set(check_status_code))):
                        
                        confidence_level = 95 if next(iter(set(check_status_code))) == 'PLA' else 90
                        confidence_level = 80 if next(iter(set(check_st_vntge_yr))) == vintage_yr and confidence_level != 95 else confidence_level

                        path_list.append((i, next(iter(set(check_st_vntge_yr))), len(i),confidence_level))
                        name_ =self.filter_path_name(df, i)
                        if name_:
                            filter_by_path_name.extend(name_)
    
                if ivs == 'PDA':
                    if ivs == next(iter(set(check_status_code))) \
                        or 'IPL' == next(iter(set(check_status_code))):
                        confidence_level = 95 if next(iter(set(check_status_code))) == 'PDA' else 90
                        confidence_level = 80 if next(iter(set(check_st_vntge_yr))) == vintage_yr and confidence_level != 95 else confidence_level

                        path_list.append((i, next(iter(set(check_st_vntge_yr))), len(i),confidence_level))
                        name_ = self.filter_path_name(df, i)
                        if name_:
                            filter_by_path_name.extend(name_)
    
                if ivs == 'IPL':
                    if ivs == next(iter(set(check_status_code))):
                        path_list.append((i, next(iter(set(check_st_vntge_yr))), len(i),95))
                        name_ = self.filter_path_name(df, i)
                        if name_:
                            filter_by_path_name.extend(name_)
    
        # if len(path_list) == 1:
            # return path_list[0][0]
          
        # else:
        return_edges = [i[0] for i in path_list if i[1] == vintage_yr]
    
        if len(return_edges) == 1:
            # return return_edges[0]
            shortest_path= return_edges[0]
        elif len(return_edges) > 1:
            temp_G =  self.set_G_edges(df, filter=True, filter_by_path_name=filter_by_path_name, vintage_yr=vintage_yr)
            shortest_path = nx.shortest_path(temp_G, source=source_node, target=destination_node, weight='weight')
            # return nx.shortest_path(temp_G, source=source_node, target=destination_node, weight='weight')
        else:
            # temp_G = self.set_G_edges(df, filter=True, filter_by_path_name=filter_by_path_name, vintage_yr=None)
            # return nx.shortest_path(temp_G, source=source_node, target=destination_node, weight='weight')
            if len(filter_by_path_name)>1:
                temp_G =  self.set_G_edges(df, filter=True, filter_by_path_name=filter_by_path_name, vintage_yr=None)
                shortest_path = nx.shortest_path(temp_G, source=source_node, target=destination_node, weight='weight')
            else:
                shortest_path=all_path[0]

        # confidence_percentage = [i[-1] for i in path_list if i[0] == shortest_path][0]
        confidence_percentage = [i[-1] for i in path_list if i[0] == shortest_path]
        confidence_percentage = confidence_percentage[0] if len(confidence_percentage) > 0 else 70
        return shortest_path,  confidence_percentage

    def get_path_(self,df, i,column_val='name'):
        try:
            return df.loc[(df['in_structure'] == i[0]) & (df['out_structure'] == i[1]), column_val].iloc[0]
        except:
            return df.loc[(df['in_structure'] == i[1]) & (df['out_structure'] == i[0]), column_val].iloc[0]
    
    def get_path_name(self,df, path_edges):
        shortest_path_list = []
        path_name = []
        for i in path_edges:
            path_value = self.get_path_(df, i)
            shortest_path_list.append(i + (path_value,))
            path_name.append(path_value)
        return shortest_path_list, path_name
    
    
    def get_shortest_path(self,df, source_node, destination_node,G,ivs,vintage_yr):
        # G = self.set_G_edges(df)         
        data = {}
        if source_node == destination_node:
            return {"error": "check source_node and source_node both node are same"}
        try:
            data['shortest_path'],data['confidence_percentage'] =self.get_shortest_path_edges(G, df,source_node, destination_node,ivs,vintage_yr)
            shortest_path_length = nx.shortest_path_length(G, source=source_node, target=destination_node, weight='weight')
            path_edges = list(zip(data['shortest_path'], data['shortest_path'][1:]))
            data['shortest_path_list'], data['path_name'] = self.get_path_name(df, path_edges)
        except Exception as e:
            return {"error": str(e)}
        return data
        
    def filter_path_name(self,df, path):
        name_ = []
        for count, i in enumerate([path]):
            for j in range(0, len(i)-1):
                name_.append(self.get_path_(df, [i[j], i[j+1]]))
        return name_


    # def FinalizeCDIF(self):
    #     print(f"Finalizing CDIF by compressing and putting all files in a zip file  .")
    #     logging.info(f"Finalizing CDIF by compressing and putting all files in a zip file.")
        
    #     self.output_zip_file_name=""
    #     # self._WriteZipFile([self.packagemetadata_file_name.replace(self.rootDirectory+"\\",""), self.csv_file_name.replace(self.rootDirectory,""),  self.field_file_name.replace(self.rootDirectory,"")])
        
    #     self._WriteSummaryReport()

        
         
    #     logging.info( f" completed with  - "+ str(self.total_rec_processed) +" records | "+ self.output_zip_file_name   )
    def FinalizeCDIF(self):
        print("Finalizing CDIF by compressing and putting all files in a zip file.")
        logging.info("Finalizing CDIF by compressing and putting all files in a zip file.")

        zip_base_name = self._zip_basename_from_config()
        files_for_zip = list(dict.fromkeys(self._created_files)) 
        if not files_for_zip:
            msg = "No files captured for this run. Skipping zip creation."
            print(msg)
            logging.warning(msg)
            self._WriteSummaryReport()
            return
        self._WriteZipFile(files_for_zip, zip_base_name)
                                # Cleanup: delete the filesjust zipped
        extensions_to_delete = {".csv", ".fields"} 
        for rel_name in list(dict.fromkeys(self._created_files)):
                try:
                    path = os.path.join(self.rootDirectory, rel_name)
                    _, ext = os.path.splitext(rel_name)
                    if ext.lower() in extensions_to_delete and os.path.isfile(path):
                        os.remove(path)
                        logging.info(f"Deleted generated file: {rel_name}")
                except Exception as e:
                    logging.warning(f"Could not delete '{rel_name}': {e}")
        self._WriteSummaryReport()
        logging.info(
            f"Completed with {getattr(self, 'total_rec_processed', 0)} records | ZIP: {self.output_zip_file_name}")
        print(f"ZIP created: {self.output_zip_file_name}")

    def _GetReferenceFieldValue(self,fld,fld_indx,row,idval):

        return

    def _GetFieldValue(self,fld,fld_indx,row,idval):

        fld_nm=fld['Field_Name'].upper()
        src_nm=fld['Source'].upper()
        if src_nm!="":
            fld_nm=src_nm
        if len(fld_nm)==0:
            print(f"Invalid field name '{fld_nm}'. Terminatig the process." )
            return
        fld_value=""
        row_fld_val=""
        try:
            fld_value = row[fld_nm] 
            row_fld_val=fld_value
        except Exception:
            logging.warning(f"Field {fld_nm} not availble in source, blank value used.")
        fld_type=fld['Field_Type'].lower()
        fld_mandatory=fld['Mandatory'].lower()
        domain_val=fld['DomainValue'].lower()

        # if fld['Field_Name']=="st_rmks":
            # aa=0
        validated_obj=util_services.UTILServices.check_and_convert_value(row_fld_val, fld_type)
          
        if validated_obj['is_null']==True or validated_obj['is_nan']  ==True:
            self.fields_issues[fld_indx][1]=self.fields_issues[fld_indx][1]+1

            fld_value=""
        else:
            if validated_obj['type_match']==True:
                fld_value=  validated_obj['converted_value']
                if fld['ReplaceValue']!="":
                    fld_value = util_services.UTILServices.get_mapped_value( fld['ReplaceValue']  ,fld_value)

                if domain_val !="":
                    if util_services.UTILServices.check_domain_value(domain_val,fld_value.lower())==False:
                        self.fields_issues[fld_indx][2]=self.fields_issues[fld_indx][2]+1
                            # if value_is_as_per_type==False:
                        self.bypassed_data.append([idval,fld_nm,fld_value ])
                
                if fld_type.startswith("string"):
                    match=re.search(r'string\((\d+)\)',fld_type)
                    max_len=int(match.group(1))
                    # modified_value = str(fld_value)   
                    strlen= len(fld_value)                             
                    if max_len<strlen:
                        fld_value = fld_value[:max_len]
                        self.string_truncate_object.append([idval,fld_nm, row_fld_val,fld_value])
                        self.fields_issues[fld_indx][3]=self.fields_issues[fld_indx][3]+1
                        trunc_record=True    
            else:
                self.fields_issues[fld_indx][5]=self.fields_issues[fld_indx][5]+1
                self.type_mismatch_data.append([idval,fld_nm,fld_type,fld_value])
                fld_value=""  # Changing value to NULL to avoid import issue in IQQGeo 
                db_type_issue=True
        if fld_value=="":
            if fld["defaultvalue"]!="":
                fld_value= fld["defaultvalue"]
        if fld_nm=="ACCOUNT_CODE":
            fld_value = std_descr.get_standard_description(fld_value)

        return fld_value 
    def _ValidateFieldValue(self,fld,fld_indx,fld_value,idval):

        fld_nm=fld['Field_Name'].upper()
        src_nm=fld['Source'].upper()
        if src_nm!="":
            fld_nm=src_nm
        if len(fld_nm)==0:
            print(f"Invalid field name '{fld_nm}'. Terminatig the process." )
            return
        
        row_fld_val=""
        try:
             
            row_fld_val=fld_value
        except Exception:
            logging.warning(f"Field {fld_nm} not availble in source, blank value used.")
        fld_type=fld['Field_Type'].lower()
        fld_mandatory=fld['Mandatory'].lower()
        domain_val=fld['DomainValue'].lower()

        # if fld['Field_Name']=="st_rmks":
            # aa=0
        validated_obj=util_services.UTILServices.check_and_convert_value(row_fld_val, fld_type)
          
        if validated_obj['is_null']==True or validated_obj['is_nan']  ==True:
            self.fields_issues[fld_indx][1]=self.fields_issues[fld_indx][1]+1

            fld_value=""
        else:
            if validated_obj['type_match']==True:
                fld_value=  validated_obj['converted_value']
                if fld['ReplaceValue']!="":
                    fld_value = util_services.UTILServices.get_mapped_value( fld['ReplaceValue']  ,fld_value)

                if domain_val !="":
                    if util_services.UTILServices.check_domain_value(domain_val,fld_value.lower())==False:
                        self.fields_issues[fld_indx][2]=self.fields_issues[fld_indx][2]+1
                            # if value_is_as_per_type==False:
                        self.bypassed_data.append([idval,fld_nm,fld_value ])
                
                if fld_type.startswith("string"):
                    match=re.search(r'string\((\d+)\)',fld_type)
                    max_len=int(match.group(1))
                    # modified_value = str(fld_value)   
                    strlen= len(fld_value)                             
                    if max_len<strlen:
                        fld_value = fld_value[:max_len]
                        self.string_truncate_object.append([idval,fld_nm, row_fld_val,fld_value])
                        self.fields_issues[fld_indx][3]=self.fields_issues[fld_indx][3]+1
                        trunc_record=True    
            else:
                self.fields_issues[fld_indx][5]=self.fields_issues[fld_indx][5]+1
                self.type_mismatch_data.append([idval,fld_nm,fld_type,fld_value])
                fld_value=""  # Changing value to NULL to avoid import issue in IQQGeo 
                db_type_issue=True
        if fld_value=="":
            if fld["defaultvalue"]!="":
                fld_value= fld["defaultvalue"]
        if fld_nm=="ACCOUNT_CODE":
            fld_value = std_descr.get_standard_description(fld_value)

        return fld_value 
    
    # def _WriteZipFile(self,selected_files):
    #     self.output_zip_file_name = self.rootDirectory +"\\" + self.target_object.upper() +'_cdif'+ datetime.now().strftime("%Y%m%d_%H%M%S")+'.zip'

    #     with zipfile.ZipFile(self.output_zip_file_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
    #         for file_name in selected_files:
    #             file_path = os.path.join(self.rootDirectory, file_name)
    #             if os.path.exists(file_path) and os.path.isfile(file_path):
    #                 zipf.write(file_path, arcname=file_name) # arcname ensures only the filename is used in the zip
    #             else:
    #                 print(f"Warning: File '{file_name}' not found in '{self.folder_path}' and could not be added to the zip.")
    #                 logging.warning(f"Warning: File '{file_name}' not found in '{self.folder_path}' and could not be added to the zip.")

    def _zip_basename_from_config(self) -> str:
        return os.path.splitext(os.path.basename(self.config_json_file_path))[0]

    def _WriteZipFile(self, selected_files: list[str], zip_base_name: str):

        root_dir = os.path.abspath(self.rootDirectory)        # e.g., .../output
        zip_dir  = os.path.abspath(self.folder_path)          # e.g., .../output/equipment
        os.makedirs(zip_dir, exist_ok=True)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        zip_name = f"{zip_base_name}_cdif{ts}.zip"
        self.output_zip_file_name = os.path.join(zip_dir, zip_name)

        with zipfile.ZipFile(self.output_zip_file_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for rel_name in selected_files:
                file_path = os.path.join(root_dir, rel_name)
                if os.path.isfile(file_path):
                    #zipf.write(file_path, arcname=rel_name)
                    zipf.write(file_path, arcname=os.path.basename(rel_name))
                else:
                    msg = (
                        f"Warning: File '{rel_name}' not found under '{root_dir}' "
                        f"and could not be added to the zip."
                    )
                    print(msg)
                    #logging.warning(msg)

    def _WriteCSVFile(self,file_name,clm_name, csvdata):
        if len(csvdata)>0:
            with open(file_name,'w',newline='',encoding='utf-8') as csvfile:
                    writer=csv.writer(csvfile)
                    # for csvdata in csvdata_list:
                    writer.writerow( clm_name)        
                    writer.writerows( csvdata)
                    try:
                                rel = os.path.relpath(file_name, self.rootDirectory)
                                self._created_files.append(rel)
                    except Exception:
                                # If relpath fails for any reason, ignore rather than crashing
                                pass         
        else:
            print("No data processed to write the CSV file.")                    
    def _WriteSummaryReport(self ):
        
        df_xl1 = pd.DataFrame(self.string_truncate_object, columns=['Object ID','Field Name', 'Original String', 'Truncated String'])
        df_xl2 = pd.DataFrame(self.duplicate_geom, columns=['Object ID','Name','Geom'])
        df_xl3 = pd.DataFrame(self.bypassed_data, columns=['Object ID','Bypassed for Field','Field Value'])
        df_xl4 = pd.DataFrame(self.type_mismatch_data, columns=['Object ID','Field Name','Field Type','Field Value'])
        df_xl5 = pd.DataFrame(self.topology_issues, columns=['Object ID','Start Point Distance','End Point Distance'])
        
        if len(self.skipped_conduit)>0:
            df_xl6 = pd.DataFrame(self.skipped_conduit, columns=['FROM_STRUCTURE_NAME', 'TO_STRUCTURE_NAME', 'STRUCTURE_HASH','OBJ_ID'])
        if len(self.core_hole_merged)>0:
            df_xl7 = pd.DataFrame(self.core_hole_merged, columns=['ACTION_TAKEN','CORE HOLE ID','FROM_STRUCTURE_NAME', 'TO_STRUCTURE_NAME', 'SPAN ID GEOMETRY UPDATED'])
        if len(self.parent_missing)>0:
            df_xl8 = pd.DataFrame(self.parent_missing, columns=['Object ID','Type','From Structure','To Structure','Length of Conduit','Length of all parts'])
        
        summary_path = os.path.join(self.rootDirectory,self.folder_name+'_summary')
        if not os.path.exists(summary_path):
            os.makedirs(summary_path ) 
        excel_file_path = summary_path+'\\SummaryReport_'+ self.target_object.upper() +"_"+datetime.now().strftime("%Y%m%d_%H%M%S")+".xlsx"
        #excel_file_path = self.rootDirectory  + '\\SummaryReport_'+ self.target_object.upper() +"_"+datetime.now().strftime("%Y%m%d_%H%M%S")+".xlsx"
        # df_xl.to_excel(excel_file_path, index=False, sheet_name='TruncatedValues')

        with pd.ExcelWriter(excel_file_path, engine='xlsxwriter') as writer:
            df_xl1.to_excel(writer, sheet_name='TruncatedValues', index=False)
            df_xl2.to_excel(writer, sheet_name='DuplicateGeom', index=False)
            df_xl3.to_excel(writer, sheet_name='OutsideDomain_Bypassed', index=False)
            df_xl4.to_excel(writer, sheet_name='Missing_Mismatch', index=False)
            df_xl5.to_excel(writer, sheet_name='TopologyIssue', index=False)

            if len(self.skipped_conduit)>0:
                df_xl6.to_excel(writer, sheet_name='SkippedConduit', index=False)
            if len(self.core_hole_merged)>0:
                df_xl7.to_excel(writer, sheet_name='CoreHole', index=False)
            if len(self.parent_missing)>0:
                df_xl8.to_excel(writer, sheet_name='MissingObjects', index=False)

        end_time = time.time()
        elapsed_seconds = end_time - self.start_time

        wb = load_workbook(excel_file_path)
        # Create a new sheet
        ws_openpyxl = wb.create_sheet('Summary',0)

        # Add data to the new sheet using openpyxl
        ws_openpyxl['A1'] = 'NE to IQGeo Object Migration Summary Report : ' + self.target_object.upper()
        ws_openpyxl.merge_cells('A1:D1')
        ws_openpyxl['A1'].font = Font(name='Arial', size=14, bold=True, color="FF0000") 
        ws_openpyxl['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws_openpyxl['A2'] = 'Start Time'
        ws_openpyxl['A3'] = "End Time"
        ws_openpyxl['A4'] = "Total Time Taken (min:sec)"
        ws_openpyxl['A5'] = "NE Source FileGDB"
        ws_openpyxl['A6'] = "NE Source Feature Layer"
        ws_openpyxl['A7'] = "Total Records in NE Source Feature Layer"
        ws_openpyxl['A8'] = "Generated CDIF File"
        ws_openpyxl['A9'] = "IQGeo Object Name"
    
        ws_openpyxl['A10'] = "Condition Used to Select NE Data"
        ws_openpyxl['A11'] = "Total Records selected as per Condition"
        ws_openpyxl['A12'] = "Total Records Processed"
        ws_openpyxl['A13'] ="Records Skipped due to value / data type issue"
        ws_openpyxl['A14'] ="Records Skipped due to reference /parent data missing reason"
        ws_openpyxl['A15'] ="Number of Records where data trucated due to shorter length"
        ws_openpyxl['A16'] ="Number of values where data trucated due to shorter length"
        ws_openpyxl['A17'] ="Number of Duplicate Geometires"
        
        m, s = divmod(end_time-self.start_time, 60)

        ws_openpyxl['B2'] = datetime.fromtimestamp(self.start_time).strftime("%d-%b-%Y %H:%M:%S")
        ws_openpyxl['B3'] = datetime.fromtimestamp(end_time).strftime("%d-%b-%Y %H:%M:%S")
        ws_openpyxl['B4'] = str(int(m)) + ":" + str(int(s))
        ws_openpyxl['B5'] = self.input_gdb
        ws_openpyxl['B6'] =  self.src_lyr
        ws_openpyxl['B7'] = len(self.org_gdf)
        ws_openpyxl['B8'] = self.output_zip_file_name

        ws_openpyxl['B9'] = self.target_object
        ws_openpyxl['B10'] = self.filter_column
        ws_openpyxl['B11'] = self.total_records
        ws_openpyxl['B12'] = self.total_rec_processed
        ws_openpyxl['B13'] = self.records_skipped_due_to_issue
        ws_openpyxl['B14'] = self.records_skipped_due_to_parent_data_mismatch_issue
        ws_openpyxl['B15'] = self.number_of_records_truncated
        ws_openpyxl['B16'] =len(self.string_truncate_object)
        ws_openpyxl['B17'] =self.number_of_geometry_duplicate

        

        ws_openpyxl.merge_cells('B2:D2')
        ws_openpyxl.merge_cells('B3:D3')
        ws_openpyxl.merge_cells('B4:D4')
        ws_openpyxl.merge_cells('B5:D5')
        ws_openpyxl.merge_cells('B6:D6')
        ws_openpyxl.merge_cells('B7:D7')
        ws_openpyxl.merge_cells('B8:D8')
        ws_openpyxl.merge_cells('B9:D9')
        ws_openpyxl.merge_cells('B10:D10')
        ws_openpyxl.merge_cells('B11:D11')
        ws_openpyxl.merge_cells('B12:D12')
        ws_openpyxl.merge_cells('B13:D13')
        ws_openpyxl.merge_cells('B14:D14')
        ws_openpyxl.merge_cells('B15:D15')
        ws_openpyxl.merge_cells('B16:D16')
        
        
        ws_openpyxl['A18'] ="Field Wise Abstract"
        ws_openpyxl.merge_cells('A18:D18')
        ws_openpyxl['A18'].font = Font(bold=True)
        ws_openpyxl['A18'].alignment  = Alignment(horizontal='center', vertical='center')

        ws_openpyxl['A19'] = 'Field Name'
        ws_openpyxl['B19'] = 'Mandatory ?'
        ws_openpyxl['C19'] = 'NULL Values Found'
        ws_openpyxl['D20'] = 'Migrated with Issues'
    
        ws_openpyxl['D20'] = 'Outside Domain/Pick List'
        ws_openpyxl['E20'] = 'Data Truncated'
        ws_openpyxl['F19'] = 'Migrated without Issue'
        ws_openpyxl['G19'] = 'Type Mismatch-Not Migrated'
        
        ws_openpyxl.merge_cells('A19:A20')
        ws_openpyxl.merge_cells('B19:B20')
        ws_openpyxl.merge_cells('C19:C20')
        ws_openpyxl.merge_cells('D19:E19')
        ws_openpyxl.merge_cells('F19:F20')
        ws_openpyxl.merge_cells('G19:G20')





        ws_openpyxl.column_dimensions['A'].width = 60
        ws_openpyxl.column_dimensions['B'].width = 20
        ws_openpyxl.column_dimensions['C'].width = 30
        ws_openpyxl.column_dimensions['D'].width = 20
        ws_openpyxl.column_dimensions['E'].width = 20
        ws_openpyxl.column_dimensions['E'].width = 20
        ws_openpyxl.column_dimensions['E'].width = 20
        
        rwno=21
        for fld in self.fields_issues:
            null_values = fld[1] if self.total_rec_processed >= fld[1] else self.total_rec_processed 
            ws_openpyxl['A'+str(rwno)] = fld[0]
            ws_openpyxl['B'+str(rwno)] = fld[4]
            ws_openpyxl['C'+str(rwno)] = null_values
            ws_openpyxl['D'+str(rwno)] = fld[2] 
            ws_openpyxl['E'+str(rwno)] = fld[3]
            ws_openpyxl['F'+str(rwno)] = self.total_rec_processed-null_values
            ws_openpyxl['G'+str(rwno)] = fld[5]
            if  fld[4].lower()=="yes":
                if int(fld[1])>0:
                    ws_openpyxl['C'+str(rwno)].font = Font(bold=True, color="FF0000") 
            
            if int(fld[2])>0:
                    ws_openpyxl['D'+str(rwno)].fill = PatternFill(start_color="FFFFCC00", end_color="FFFFCC00", fill_type="solid") # Goldenrod yellow

            if int(fld[3])>0:
                    ws_openpyxl['E'+str(rwno)].fill = PatternFill(start_color="ff6900", end_color="ff6900", fill_type="solid") 

            rwno+=1
        # Save the modified workbook
        wb.active = ws_openpyxl
        wb.save(excel_file_path)
     
    def _UpdateMessengerGeom(self,multiple_msngr, csv_data):
        print("Updating MESSANGER Geometry")
        # query = f"select obj_id, from_str_id, to_str_id , str_hash,geom from route_geom rg  where str_hash in (select  str_hash from route_geom rg  where span_type like 'MESS%'  group by str_hash having count(1)>1) order by str_hash, no_vertex "
        shortest_geom_obj_id=''
        shortest_geom= ''
        last_str_hash=''
        for crw in multiple_msngr:
            cur_str_hash=crw[3]
            if last_str_hash!=cur_str_hash:
                #bypassing 1st row of group which has lowest vertices / length
                # shortest_geom_obj_id=crw[0]
                shortest_geom=crw[4]
            else:
                for row in csv_data:
                    if row[0] == crw[0]  :
                        row[1]=shortest_geom
                        break
            last_str_hash=cur_str_hash
        return csv_data
            
    def _UpdateBuriedSpanGeomByMergeCoreHole(self,csv_data,insert_vertices_at,ch_span_geom , buried_span_id_to_update_geom,ch_span_id ,FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME):
        
        geom_updated=False
        for row in csv_data:
            if row[0] == buried_span_id_to_update_geom:  
                buried_span_geom= row[1]
                print(f"Merging CoreHole Geometry for- {buried_span_id_to_update_geom} from geom {ch_span_geom} to geom {buried_span_geom} ")

                # updated_geom = util_services.UTILServices.insert_geometry(buried_span_geom,ch_span_geom,insert_vertices_at)
                updated_geom = util_services.UTILServices.merge_multilines_wkb(buried_span_geom,ch_span_geom)
                updated_geom= updated_geom.upper()
                row[1] = updated_geom
                self.core_hole_merged.append(['CORE HOLE GEOM MERGED',ch_span_id,FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME , buried_span_id_to_update_geom])
                geom_updated=True
                break
            
        if geom_updated==False:
            self.core_hole_merged.append(['CORE HOLE GEOM COULD NOT MERGED',ch_span_id,FROM_STRUCTURE_NAME,TO_STRUCTURE_NAME , buried_span_id_to_update_geom])
            
        return csv_data
    
    def _CreateConduitRun(self,conduit_data ,clm_name,span_name_index):
        # self._WriteCSVFile("test_conduit.csv",clm_name,conduit_data)
        print('Processing conduit_run...')
        field_file_name= self.folder_path+"\\mywcom_conduit_run.fields"
        # clmcsv=io.StringIO()
        # writer = csv.writer(clmcsv)
        # writer.writerow(clm_name)
        # csvstring =clmcsv.getvalue()
        conduit_conduitrun_data=[]
        cr_data=[]
        # try:
        hdr_info=[["name","type","unit"]]
        hdr_info.append(["id","id",""]) #ID field column information 
        hdr_info.append(["path","path",""])  
        with open(field_file_name,'w',newline='') as csvfile:
            writer=csv.writer(csvfile)
            writer.writerows(hdr_info)
        
        print(f"Preparing Conduit Data for processing...")

        conduit_df = pd.DataFrame(conduit_data)

        

        print(f"Querying database for span units associated with multiple span...")
        
        
        # params = oracledb.ConnectParams(host="sod-oradb-023.stholdco.com", port=1527, service_name="DBD463.STHOLDCO.COM")
        # params = oracledb.ConnectParams(host=dbcon['host'], port=dbcon['port'], service_name=dbcon['service'])
        # ora_conn2 = oracledb.connect(user=dbcon['user'], password=dbcon['password'], params=params)
        # ora_cursor2 = ora_conn2.cursor()
        sql_conn = sqlite3.connect(self.sqlite_db)
        sql_cursor = sql_conn.cursor()
        # sql_query = f"SELECT mssu.span_unit_uuid FROM ne.MV_SPAN_SPAN_UNIT mssu INNER JOIN ne.MV_SPAN ms ON mssu.SPAN_UUID = ms.UUID GROUP BY span_unit_uuid HAVING count(DISTINCT span_uuid ) >1"
        sql_query = f"SELECT mssu.span_unit_uuid FROM SPAN_SPAN_UNIT mssu INNER JOIN SPAN ms ON mssu.SPAN_UUID = ms.UUID GROUP BY span_unit_uuid HAVING count(DISTINCT span_uuid ) >1"
        # sql_query = f"SELECT mssu.span_unit_uuid FROM SPAN_SPAN_UNIT mssu INNER JOIN SPAN ms ON mssu.SPAN_UUID = ms.UUID GROUP BY span_unit_uuid "

        # ora_cursor2.execute(sql_query)
        # rows = ora_cursor2.fetchall()
        sql_cursor.execute(sql_query)
        rows = sql_cursor.fetchall()
        rw_no=0
        print(f"Total {len(rows)} found.")
        geom_col_index=util_services.UTILServices.GetColumnIndex(clm_name, "path")
        for span_unit_row in rows:
            fld_indx=0
            geom_ary=[]
            csv_row=[]
            span_span_uuid=span_unit_row[0]     
            rw_no+=1
            # conduit_run_id="CR" + re.sub(r'[^0-9a-zA-Z]', '', span_span_uuid)   
            conduit_run_id="mywcom_conduit_run/" + str(rw_no)   
            
            # print(f"Processing for {span_span_uuid}")
            #select all the span associated so that there geometry can be mereged to create conduit run
            # sql_query = f"SELECT ms.SPAN_NAME FROM ne.MV_SPAN_SPAN_UNIT mssu INNER JOIN ne.MV_SPAN ms ON mssu.SPAN_UUID = ms.UUID  WHERE mssu.SPAN_UNIT_UUID = '{span_span_uuid}'"
            sql_query = f"SELECT ms.SPAN_NAME FROM SPAN_SPAN_UNIT mssu INNER JOIN SPAN ms ON mssu.SPAN_UUID = ms.UUID  WHERE mssu.SPAN_UNIT_UUID = '{span_span_uuid}'"
            # ora_cursor2.execute(sql_query)
            print(sql_query)
            sql_cursor.execute(sql_query)
            # associated_span_rows = ora_cursor2.fetchall()
            # associated_span= [r[0] for r in ora_cursor2.fetchall()]
            associated_span= [r[0] for r in sql_cursor.fetchall()]

            # filtered_df = conduit_df[conduit_df.iloc[:,span_name_index].isin(associated_span)]

            filtered_df = conduit_df[conduit_df[span_name_index].isin(associated_span)]
            # print(len(filtered_df))
            if (len(filtered_df)>1):
                for index,rw in filtered_df.iterrows():
                    geom_ary.append(rw[geom_col_index])
                    conduit_conduitrun_data.append([conduit_run_id,rw[0],rw[span_name_index]])
                geom_merged=util_services.UTILServices.merge_geom(geom_ary)
                csv_row.append(conduit_run_id)
                csv_row.append(geom_merged)
                cr_data.append(csv_row)
            else:
                if (len(filtered_df)==1):
                    self.parent_missing.append([span_span_uuid,"MISMATCH BETWEEN DB AND NE DATA", 'SINGLE SPAN FOUND IN SPAN TABLE' , ''])
                    
                    # csv_row.append(conduit_run_id)
                    # csv_row.append(filtered_df.iloc[0,geom_col_index]  )
                    # cr_data.append(csv_row)
                    print(f'Single geom found for span_unit_uuid {span_span_uuid} in NE Data which mismatch with DB data hence skipping')
                else:
                    print(f'No geometry found to merge for the span_unit_uuid {span_span_uuid}')
                    

        # ora_cursor2.close
        # ora_conn2.close
        sql_cursor.close()
        sql_conn.close()
        # conduit_run_file_csv=self.folder_path+"\\mywcom_conduit_run.csv"
        # print(f'Writing conduit_run csv file.{conduit_run_file_csv}')
        # self._WriteCSVFile(conduit_run_file_csv ,['id','path'],cr_data)        
        return conduit_conduitrun_data,cr_data

        # except Exception as e:
        #     print(f'Conduit Run field metadata file could not be written. Error occured : {str(e)}')
        #     logging.error( f"Error occured in create Conduit Run field mapping file '{field_file_name}'. Error is {str(e)}  ")
        #     return